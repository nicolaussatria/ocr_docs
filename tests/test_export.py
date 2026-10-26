"""The claim forms as one Excel workbook (no OCR models needed)."""
import base64
import io
from datetime import date

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image, ImageDraw
from pydantic import ValidationError

from docscan.export import (ATTACH_SHEET, CLAIM_SHEET, MEDICAL_SHEET, ClaimExport, ClaimRow, ClaimTable, MedicalRow,
                            build_xlsx, medical_claim)


def picture(w=300, h=700, text="STRUK") -> str:
    """A made-up receipt picture, base64 JPEG like the scan returns (never a real receipt in a test)."""
    im = Image.new("RGB", (w, h), "white")
    ImageDraw.Draw(im).text((20, 20), text, fill="black")
    buf = io.BytesIO()
    im.save(buf, format="JPEG")
    return base64.b64encode(buf.getvalue()).decode()

ROWS = [ClaimRow(date=date(2026, 9, 28), liters=18.8, bbm=188000, note="BBM Pertamax 18,8 L - SPBU Cipularang B"),
        ClaimRow(date=date(2026, 9, 28), tol=11000, note="Tol"),
        ClaimRow(liters=3.5, bbm=35000, note="BBM Pertalite 3,5 L (bon kontan) (tanggal tidak terbaca)")]
MEDICAL = [MedicalRow(date=date(2026, 9, 2), amount=700000, service="Konsultasi, Amoxicillin 500mg", patient="Andi Setiawan",
                      note="Klinik Pratama Sehat - Jl. Raya Bogor No. 12, Jakarta Timur - dr. Budi Santoso"),
           MedicalRow(amount=35500, service="Paracetamol 500mg", note="Apotek Kimia Farma 312")]


def book(claim: ClaimExport):
    return load_workbook(io.BytesIO(build_xlsx(claim)))


def merged(ws):
    return {str(m) for m in ws.merged_cells.ranges}


def test_the_sheet_looks_like_the_paper_form():
    ws = book(ClaimExport(tables=[ClaimTable(name="page-1-fuel.jpg", rows=ROWS)]))[CLAIM_SHEET]
    assert [ws[f"{c}3"].value for c in "ABCFGHIJKL"] == ["No", "Tanggal", "BBM", "Tol", "Parkir", "Pulsa",
                                                          "Servis Kendaraan", "Lain-lain", "Total", "Keterangan"]
    assert [ws[f"{c}4"].value for c in "CDE"] == ["KM", "Liter", "Rp"]
    assert {"C3:E3", "A3:A4", "A1:L1"} <= merged(ws)
    assert ws["A1"].value == "Claim Tunjangan Operasional Karyawan" and ws["A1"].font.bold
    assert ws["A2"].value is None and ws.print_title_rows == "$1:$4"         # one claim: no name above its table
    # one row per receipt - two on the same day stay two rows
    assert [ws[f"A{r}"].value for r in (5, 6, 7)] == [1, 2, 3]
    assert ws["B5"].value.date() == date(2026, 9, 28) and ws["B5"].number_format == "dd/mm/yyyy"
    assert ws["B7"].value is None                                    # no readable date: left for the reviewer
    assert ws["D5"].value == 18.8 and ws["D5"].number_format == "0.00"
    assert (ws["E5"].value, ws["F6"].value, ws["K5"].value, ws["K6"].value) == (188000, 11000, 188000, 11000)
    assert ws["E5"].number_format == "#,##0" and ws["F5"].value is None   # nothing claimed there: empty, not 0
    assert ws["L7"].value.endswith("(tanggal tidak terbaca)")
    # the Total row
    assert ws["B8"].value == "Total" and ws["D8"].value == 22.3 and ws["E8"].value == 223000 and ws["K8"].value == 234000
    assert ws["K8"].font.bold and ws.freeze_panes == "A5"


def test_claims_per_document_are_stacked_under_their_names():
    wb = book(ClaimExport(tables=[ClaimTable(name="page-1-fuel.jpg", rows=ROWS[:2]),
                                  ClaimTable(name="Kwitansi-2.png", rows=ROWS[2:])]))
    assert wb.sheetnames == [CLAIM_SHEET]                            # one sheet, as on paper
    ws = wb[CLAIM_SHEET]
    assert ws["A2"].value == "page-1-fuel" and ws["A2"].font.italic
    assert ws["A3"].value == "No" and ws["B7"].value == "Total" and ws["K7"].value == 199000
    # two empty rows, then the next document's name, headings and its own total
    assert ws["A8"].value is None and ws["A9"].value is None and ws["A10"].value == "Kwitansi-2"
    assert ws["A11"].value == "No" and ws["C12"].value == "KM" and {"C11:E11", "A11:A12"} <= merged(ws)
    assert ws["A13"].value == 1 and ws["E13"].value == 35000 and ws["B14"].value == "Total" and ws["K14"].value == 35000


def test_medical_receipts_get_their_own_claim_form():
    wb = book(ClaimExport(tables=[ClaimTable(name="a.jpg", rows=ROWS)], medical=MEDICAL))
    assert wb.sheetnames == [CLAIM_SHEET, MEDICAL_SHEET]
    ws = wb[MEDICAL_SHEET]
    assert ws["A1"].value == "Claim Medical Karyawan" and "A1:G1" in merged(ws)
    assert [ws[f"{c}3"].value for c in "ABCDEFG"] == ["No", "Tanggal", "Nilai Kwitansi", "Nilai Claim",
                                                       "Deskripsi Pelayanan/Nama Obat", "Nama Pasien di Kwitansi", "Keterangan"]
    assert [ws[f"{c}4"].value for c in "ACDEF"] == [1, 700000, 630000, "Konsultasi, Amoxicillin 500mg", "Andi Setiawan"]
    assert ws["G4"].value.startswith("Klinik Pratama Sehat - Jl. Raya Bogor") and ws["B4"].number_format == "dd/mm/yyyy"
    assert (ws["C5"].value, ws["D5"].value, ws["F5"].value) == (35500, 31950, None)   # no patient on a pharmacy slip
    assert ws["B6"].value == "Total" and ws["D6"].value == 661950 and ws["D6"].font.bold and ws["C6"].value is None


def test_ninety_percent_in_whole_rupiah():
    assert [medical_claim(v) for v in (700000, 35500, 1005, 0)] == [630000, 31950, 905, 0]


def test_only_medical_receipts_give_only_the_medical_form():
    assert book(ClaimExport(tables=[ClaimTable(name="x", rows=[])], medical=MEDICAL[:1])).sheetnames == [MEDICAL_SHEET]


def test_export_endpoint(monkeypatch):
    monkeypatch.setenv("DOCSCAN_API_KEY", "secret")
    from docscan.api import app

    c = TestClient(app)                                                 # no OCR needed: models are not loaded
    body = {"tables": [{"name": "Claim", "rows": [{"date": "2026-09-28", "liters": 18.8, "bbm": 188000, "note": "BBM"}]}],
            "medical": [{"date": "2026-09-02", "amount": 700000, "patient": "Andi Setiawan"}]}
    assert c.post("/v1/export/xlsx", json=body).status_code == 401
    r = c.post("/v1/export/xlsx", json=body, headers={"X-API-Key": "secret"})
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
    wb = load_workbook(io.BytesIO(r.content))
    assert wb[CLAIM_SHEET]["E5"].value == 188000 and wb[MEDICAL_SHEET]["D4"].value == 630000
    assert c.post("/v1/export/xlsx", json={"tables": [], "medical": []}, headers={"X-API-Key": "secret"}).status_code == 422


def anchors(ws):
    return sorted(img.anchor._from.row + 1 for img in ws._images)    # 1-based rows the pictures start on


def test_receipt_pictures_follow_their_rows_on_an_attachment_sheet():
    rows = [r.model_copy(update={"image": picture(text=f"R{i}")}) for i, r in enumerate(ROWS)]
    med = [MEDICAL[0].model_copy(update={"image": picture(900, 400)}), MEDICAL[1]]          # one without a picture
    wb = book(ClaimExport(tables=[ClaimTable(name="page-1-fuel.jpg", rows=rows)], medical=med))
    assert wb.sheetnames == [CLAIM_SHEET, MEDICAL_SHEET, ATTACH_SHEET]
    ws = wb[ATTACH_SHEET]
    assert ws["A1"].value == "Attachment - Struk dan Kwitansi"
    headings = [ws.cell(r, 1).value for r in range(3, ws.max_row + 1) if str(ws.cell(r, 1).value).startswith("Claim")]
    assert headings == [f"{CLAIM_SHEET} - No 1", f"{CLAIM_SHEET} - No 2", f"{CLAIM_SHEET} - No 3",
                        f"{MEDICAL_SHEET} - No 1", f"{MEDICAL_SHEET} - No 2"]          # the forms' order and numbers
    assert ws["A4"].value == "28/09/2026 | BBM Rp 188.000 | 18,80 L | BBM Pertamax 18,8 L - SPBU Cipularang B"
    assert len(ws._images) == 4 and anchors(ws)[0] == 5                                  # under heading and detail
    # every picture is shown within the column of receipts, keeping its shape (the size shown is in the anchor)
    shown = sorted((round(i.anchor.ext.width / 9525), round(i.anchor.ext.height / 9525)) for i in ws._images)
    assert shown == [(274, 640), (274, 640), (274, 640), (480, 213)]
    # the medical receipt without a picture says so, under its own heading
    missing = next(r for r in range(3, ws.max_row + 1) if ws.cell(r, 1).value == f"{MEDICAL_SHEET} - No 2")
    assert ws.cell(missing + 1, 1).value.startswith("tanggal tidak terbaca | Nilai Kwitansi Rp 35.500")
    assert ws.cell(missing + 2, 1).value == "(gambar struk tidak tersedia)"
    # a row's No opens its picture, and the picture's heading goes back to the row
    claim = wb[CLAIM_SHEET]
    assert claim["A5"].hyperlink.location == f"'{ATTACH_SHEET}'!A3" and claim["A5"].value == 1
    assert ws["A3"].hyperlink.location == f"'{CLAIM_SHEET}'!A5"
    assert wb[MEDICAL_SHEET]["A5"].hyperlink.location == f"'{ATTACH_SHEET}'!A{missing}"


def test_a_receipt_is_never_cut_across_two_printed_pages():
    rows = [ClaimRow(date=date(2026, 9, d), tol=11000, image=picture(300, 640)) for d in range(1, 6)]
    ws = book(ClaimExport(tables=[ClaimTable(name="x", rows=rows)]))[ATTACH_SHEET]
    breaks = [b.id for b in ws.row_breaks.brk]
    assert len(breaks) == 4                                       # one tall receipt per page
    assert all(any(b == a - 3 for a in anchors(ws)) for b in breaks)   # each break falls right before a heading


def test_without_pictures_there_is_no_attachment_sheet():
    assert ATTACH_SHEET not in book(ClaimExport(tables=[ClaimTable(name="x", rows=ROWS)], medical=MEDICAL)).sheetnames


def test_only_pictures_are_accepted():
    with pytest.raises(ValidationError):
        ClaimRow(tol=1, image="not base64 at all!")
    with pytest.raises(ValidationError):
        ClaimRow(tol=1, image=base64.b64encode(b"GIF89a" + bytes(40)).decode())
    with pytest.raises(ValidationError):
        MedicalRow(amount=1, image=base64.b64encode(b"<svg onload=alert(1)>").decode())
    assert ClaimRow(tol=1, image="").image is None


def test_export_endpoint_with_pictures(monkeypatch):
    monkeypatch.delenv("DOCSCAN_API_KEY", raising=False)
    from docscan.api import app

    body = {"tables": [{"name": "Claim", "rows": [{"date": "2026-09-28", "tol": 11000, "note": "Tol", "image": picture()}]}]}
    r = TestClient(app).post("/v1/export/xlsx", json=body)
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == [CLAIM_SHEET, ATTACH_SHEET] and len(wb[ATTACH_SHEET]._images) == 1
